# Xarajatlar xulosasi (albom uchun): qayta hisoblangan Excel'dan (LibreOffice recalc) qiymatlarni JSON qilib chiqaradi.
#   python3 xarajat-xulosa.py → stdout (JSON)
import json
from pathlib import Path
from openpyxl import load_workbook

BU = Path(__file__).resolve().parent
fayl = sorted((BU.parent / 'hisob').glob('Xadra_xarajatlar_v*.xlsx'))[-1]
wb = load_workbook(fayl, data_only=True)
wy, ws = wb['Yakuniy'], wb["Xonalar bo'yicha"]

turlar, bolimlar, qoshimcha, bolim = [], [], {}, None
for row in wy.iter_rows(min_row=4, values_only=True):
    a, b, c, d, e = (list(row) + [None] * 5)[:5]
    if a == "2. Qavat va bo'limlar bo'yicha": bolim = True; continue
    if isinstance(a, int) and isinstance(c, (int, float)):
        (bolimlar if bolim else turlar).append({'nom': b, 'usd': c, 'som': d, 'ulush': e})
    elif b in ('JAMI', 'UMUMIY JAMI') and not bolim:
        qoshimcha[b] = {'usd': c, 'som': d}
    elif isinstance(b, str) and b.startswith('Kutilmagan'):
        qoshimcha['zaxira'] = {'usd': c, 'som': d, 'nom': b}
    elif isinstance(b, str) and b.startswith('shundan taxminiy'):
        qoshimcha['taxmin'] = {'usd': c, 'som': d}
xonalar, joriy = [], None
oxirgi = ws.max_column
for row in ws.iter_rows(min_row=5, values_only=True):
    a, jami = row[0], row[oxirgi - 1]
    if not a: continue
    if jami is None and not str(a).startswith(('JAMI', 'UMUMIY', 'Kutilmagan', 'Tekshiruv')):
        joriy = a; continue
    if str(a).endswith('— jami') or str(a).startswith(('JAMI', 'UMUMIY', 'Kutilmagan', 'Tekshiruv')): continue
    xonalar.append({'bolim': joriy, 'nom': a, 'usd': jami})
kurs = wb['Narxlar']
kurs_q = next(r[2] for r in kurs.iter_rows(min_row=5, values_only=True) if r[1] == 'Dollar kursi')
# koworking sahnasi (taxminiy narxlar, 1 tadan)
QISQA = [('Koworking sahnasi', 'karkas va grafit panel'), ('Sahna: ', 'reykalar'), ('Sahna logotipi', 'logotip'),
         ('Sahna yoritishi', 'trek yoritish va LED'), ('Sahna ekrani', "ekranni ko'chirish"), ('Ovoz:', 'ovoz tizimi'),
         ('Sahna maydoni poli', 'sahna poli'), ('Minbar', 'minbar'), ("Katta o'simlik", "o'simlik")]
sahna = [{'nom': q, 'usd': r[2]} for r in kurs.iter_rows(min_row=5, values_only=True) for p, q in QISQA
         if isinstance(r[1], str) and r[1].startswith(p) and isinstance(r[2], (int, float))]
# fazalar (Fazalar varag'i): 1-jadval, oylik pul oqimi, 3-jadval (har fazadagi xarajatlar)
wf = wb['Fazalar']
fazalar, oylar, taqsim, bosqich, sar3 = [], [], [], None, None
for row in wf.iter_rows(min_row=4, values_only=True):
    a, b = row[0], row[1]
    if isinstance(a, str) and a.startswith('1. '): bosqich = 1; continue
    if isinstance(a, str) and a.startswith('2. '): bosqich = 2; continue
    if isinstance(a, str) and a.startswith('Oylik pul'): bosqich = 'oy'; continue
    if isinstance(a, str) and a.startswith('3. '): bosqich = 3; continue
    if bosqich == 1 and isinstance(a, str) and a.startswith('F') and len(a) == 2:
        fazalar.append({'kod': a, 'nom': b, 'ish': row[2], 'kun': row[3], 'bosh': row[4].strftime('%d.%m.%Y'), 'tugash': row[5].strftime('%d.%m.%Y'),
                        'xarajat': row[6], 'zaxira': row[7], 'usd': row[8], 'som': row[9], 'ulush': row[10], 'yigma': row[11]})
    if bosqich == 'oy' and a == 'Faza': oylar = [{'oy': c.strftime('%m.%Y'), 'usd': 0} for c in row[2:6]]
    if bosqich == 'oy' and b == 'Jami, $':
        for o, v in zip(oylar, row[2:6]): o['usd'] = v
    if bosqich == 'oy' and b == "Jami, so'm":
        for o, v in zip(oylar, row[2:6]): o['som'] = v
    if bosqich == 3 and a == 'Kalit': sar3 = row; continue
    if bosqich == 3 and sar3 and a and b and a != 'Kalit' and isinstance(row[3], (int, float)):
        n = len(fazalar)
        taqsim.append({'kalit': a, 'nom': b, 'tur': row[2], 'usd': row[3], 'ulush': list(row[4:4 + n]), 'faza': list(row[5 + n:5 + 2 * n])})
print(json.dumps({'fayl': fayl.name, 'turlar': turlar, 'bolimlar': bolimlar, 'xonalar': xonalar, 'kurs': kurs_q, **qoshimcha,
                  'sahna': {'jami': sum(q['usd'] for q in sahna), 'qatorlar': sahna},
                  'fazalar': fazalar, 'oylar': oylar, 'taqsim': taqsim}, ensure_ascii=False))

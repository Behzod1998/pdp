# Xarajatlar jadvali (Excel): chizmadagi yangi devorlar, shisha, bambuk panel, interaktiv doska, perila,
# shisha eshiklar, yoritish, videokuzatuv, logotiplar, matli plyonka, yong'in xavfsizligi, buzish ishlari va
# 10% kutilmagan xarajat; 1-qavat (resepshn, logotiplar), sotuv va admin xonalari mebeli, koridor doskalari — taxminiy narxlar.
# So'mdagi narxlar «Narxlar» varag'idagi kurs bo'yicha dollarga o'tkaziladi.
# Uzunliklar model.mjs dan (xarajat-malumot.mjs orqali), narx va balandliklar «Narxlar» varag'ida —
# ularni Excelda o'zgartirsa, butun hisob formulalar bilan qayta hisoblanadi.
#   python3 xarajatlar.py            → ../hisob/Xadra_xarajatlar_v1.1.xlsx (1, 2 va 3-qavat)
import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BU = Path(__file__).resolve().parent
CHIQISH = BU.parent / 'hisob' / 'Xadra_xarajatlar_v1.1.xlsx'

d = json.loads(subprocess.check_output(['node', str(BU / 'xarajat-malumot.mjs')], text=True))

XONALAR = [
    ('Q1', "1-qavat — kirish zali (resepshn, turniket)"),
    ('SR1', "1-xona"), ('SR2', "2-xona"), ('SR3', "3-xona"), ('SR4', "4-xona"),
    ('SR5', "5-xona"), ('SR6', "6-xona"), ('SR7', "7-xona"),
    ('XZ1', "Offline sotuv (XZ1)"), ('XZ2', "Admin (XZ2)"), ('XZ3', "Ustozlar xonasi (XZ3)"), ('XZ4', "Call-markaz (XZ4)"),
    ('X12', "CEO xonasi (12-xona)"),
    ('K1', "K1 koridori"), ('K2', "K2 koridori"), ('KW', "Koworking zali"), ('ZL', "Kirish zali"),
    ('ZINA', "Zinapoya"), ('BQ', "Butun qavat (umumiy ishlar)"), ('YQ', "3-qavat 5-xona — o'quv xonasi (24 o'rin)"),
    ('HJ', "Hojatxonalar (3-qavat erkaklar, 2-qavat ayollar)"),
]
NOM = dict(XONALAR)
BOLIM = [("1-qavat: kirish zali", ['Q1']),
         ("2-qavat: o'quv xonalari", ['SR1', 'SR2', 'SR3', 'SR4', 'SR5', 'SR6', 'SR7']),
         ("2-qavat: ofis xonalari", ['XZ1', 'XZ2', 'XZ3', 'XZ4', 'X12']),
         ("2-qavat: umumiy joylar", ['K1', 'K2', 'KW', 'ZL', 'ZINA', 'BQ']),
         ("3-qavat va hojatxonalar", ['YQ', 'HJ'])]
KORIDOR = {'XZ1': 'K2', 'XZ2': 'K2', 'XZ3': 'K1', 'XZ4': 'K1'}
UMUMIY = {'K1', 'K2', 'KW', 'ZL'}

# ---------- Narxlar varag'i: kataklar ----------
N = {}  # nom → katak manzili
NARX_QATORLAR = [
    ('shisha', "Shisha devor", 55, "$/m²", "Buyurtmachi"),
    ('gkl', "Gips karton devor (bir tomoni)", 12, "$/m²", "Buyurtmachi; devorning har tomoni alohida hisoblanadi"),
    ('bambuk', "Bambuk panel", 15.5, "$/m²", "Buyurtmachi"),
    ('doska', "Interaktiv doska", 1032, "$/dona", "Buyurtmachi; 7 ta shu qavat sinflariga + 1 ta yuqori qavatdagi sinfga = 8 ta"),
    ('perila', "Perila", 84, "$/m", "Buyurtmachi"),
    ('perilaUz', "Perila uzunligi", 24, "m", "Buyurtmachi"),
    ('yorKatta', "Yoritish — koridor va koworking (har biriga)", 336, "$/joy", "Buyurtmachi: K1, K2 koridorlari va koworking zali"),
    ('yorXona', "Yoritish — xona (har biriga)", 168, "$/xona", "Buyurtmachi: 7 ta o'quv xonasi, 4 ta ofis xonasi va CEO xonasi"),
    ('kamera', "Videokamera", 40, "$/dona", "Buyurtmachi"),
    ('kameraSoni', "Videokameralar soni", 16, "dona", "Buyurtmachi"),
    ('kameraOrn', "Kamera o'rnatish (har biri)", 21, "$/dona", "Buyurtmachi"),
    ('nvr', "NVR (videoregistrator)", 840, "$/dona", "Buyurtmachi"),
    ('kabel', "Kamera kabellari (jami)", 303, "$/komplekt", "Buyurtmachi: jami summa"),
    ('logo', "Logotiplar — 2-qavat (jami)", 2500, "$/komplekt", "Buyurtmachi: jami summa. Koridor oxiridagi 2 ta yorituvchi logotip shu summaga kirgan deb olindi — alohida bo'lsa, qo'shing"),
    ('kurs', "Dollar kursi", 11900, "so'm/$", "Buyurtmachi; so'mdagi narxlar shu kurs bo'yicha $ ga o'tadi"),
    ('plyonka', "Matli plyonka (shisha devorga)", 800000, "so'm/dona", "Buyurtmachi"),
    ('plyonkaSoni', "Matli plyonka soni", 8, "dona", "Buyurtmachi"),
    ('yongin', "Yong'in xavfsizligi (signalizatsiya, o't o'chirgich, belgilar, avariya chiroqlari)", 18000000, "so'm/komplekt", "Buyurtmachi: jami summa"),
    ('buzish', "Devor buzish", 50000, "so'm/m²", "Buyurtmachi"),
    ('chiqindi', "Qurilish chiqindisini olib chiqish", 1200000, "so'm/komplekt", "Buyurtmachi: jami summa"),
    ('chashagen', "Chashagen", 700000, "so'm/dona", "Buyurtmachi"),
    ('chashagenSoni', "Chashagen soni", 4, "dona", "Buyurtmachi: 3 ta erkaklar (3-qavat), 1 ta ayollar (2-qavat)"),
    ('unitaz', "Unitaz", 1200000, "so'm/dona", "Buyurtmachi"),
    ('unitazSoni', "Unitaz soni", 3, "dona", "Buyurtmachi: 1 ta erkaklar, 2 ta ayollar"),
    ('rakovina', "Rakovina", 650000, "so'm/dona", "Buyurtmachi"),
    ('rakovinaSoni', "Rakovina soni", 4, "dona", "Buyurtmachi: 2 ta erkaklar hojatxonasida, qolgan 2 tasi ayollarnikida"),
    ('mustahab', "Mustahab", 200000, "so'm/dona", "Buyurtmachi"),
    ('mustahabSoni', "Mustahab soni", 7, "dona", "Buyurtmachi: har bir unitaz va chashagenga (4 + 3)"),
    ('taxoratSm', "Taxorat smesiteli", 400000, "so'm/dona", "Buyurtmachi"),
    ('taxoratSmSoni', "Taxorat smesiteli soni", 4, "dona", "Buyurtmachi: 3 tasi erkaklar taxorat joyida, 1 tasi ayollarnikida"),
    ('rakovinaSm', "Rakovina smesiteli", 400000, "so'm/dona", "Buyurtmachi"),
    ('rakovinaSmSoni', "Rakovina smesiteli soni", 4, "dona", "Buyurtmachi: har bir rakovinaga"),
    ('tochka', "Santexnika nuqtasi (usta ishi)", 1000000, "so'm/nuqta", "Buyurtmachi"),
    ('tochkaSoni', "Santexnika nuqtalari soni", 16, "nuqta", "Buyurtmachi"),
    ('kafel', "Kafel (material)", 100000, "so'm/m²", "Buyurtmachi"),
    ('kafelUsta', "Kafel yotqizish (usta)", 100000, "so'm/m²", "Buyurtmachi"),
    ('kafelDevor', "Kafel — erkaklar hojatxonasi devori", 58, "m²", "Buyurtmachi"),
    ('kafelPol', "Kafel — erkaklar hojatxonasi poli", 16, "m²", "Buyurtmachi"),
    ('kabina', "Hojatxona kabinasi", 1200000, "so'm/dona", "Buyurtmachi"),
    ('kabinaSoni', "Hojatxona kabinalari soni", 4, "dona", "Buyurtmachi: erkaklar hojatxonasi (3-qavat)"),
    ('issiqQuvur', "Issiq suv quvuri", 25000, "so'm/m", "Buyurtmachi"),
    ('issiqUz', "Issiq suv quvuri uzunligi", 40, "m", "Buyurtmachi"),
    ('sovuqQuvur', "Sovuq suv quvuri", 12000, "so'm/m", "Buyurtmachi «12 0000» deb yozgan — 12 000 deb olindi (issiq suv quvuridan arzon); tekshiring"),
    ('sovuqUz', "Sovuq suv quvuri uzunligi", 40, "m", "Buyurtmachi"),
    ('aksessuar', "Quvur aksessuarlari (fitinglar)", 5000, "so'm/dona", "Buyurtmachi"),
    ('aksessuarSoni', "Quvur aksessuarlari soni", 100, "dona", "Buyurtmachi: jami 500 000 so'm"),
    ('kanal', "Kanalizatsiya quvuri (40 m) va otvodlar", 4000000, "so'm/komplekt", "Buyurtmachi: jami summa"),
    ('kond2', "Konditsioner — 2-qavat (2 ta, 100 mingtalik, o'rnatish bilan)", 7000, "$/komplekt", "Buyurtmachi: 2 ta uchun jami, o'rnatish bilan"),
    ('vozdux', "Havo kanallari (vozduxovod), 200 m²", 7000, "$/komplekt", "Buyurtmachi: jami summa"),
    ('kond3', "Konditsioner — 3-qavat o'quv xonasi", 1000, "$/dona", "Buyurtmachi"),
    # Taxminiy narxlar (buyurtmachi: «narxi va qanday mebel bo'lishi noma'lum — taxminan narx qo'yib berasan»)
    ('stolOfis', "Ofis stoli 140 × 70 (LDSP, tumba bilan)", 220, "$/dona", "TAXMIN — Toshkent bozori; mebelchi narxi bilan almashtiring"),
    ('kreslo', "Ofis kreslosi (to'r suyanchiqli)", 130, "$/dona", "TAXMIN"),
    ('mijozStul', "Mijoz / mehmon stuli", 60, "$/dona", "TAXMIN"),
    ('divan', "Divan, 2–3 o'rinli (kutish joyi)", 450, "$/dona", "TAXMIN"),
    ('shkaf', "Hujjat shkafi / stellaj", 350, "$/dona", "TAXMIN"),
    ('stolcha', "Jurnal stolchasi", 100, "$/dona", "TAXMIN"),
    ('interyerXZ1', "Sotuv xonasi interyeri: brend devor (logotip, plakat), o'simliklar, dekor", 600, "$/komplekt", "TAXMIN"),
    ('interyerXZ2', "Admin xonasi interyeri: dekor, o'simliklar", 300, "$/komplekt", "TAXMIN"),
    ('elonDoska', "E'lonlar doskasi 1.8 × 1.0 m (probka/magnit, alyuminiy ramka)", 120, "$/dona", "TAXMIN"),
    ('etirofDoska', "E'tirof doskasi 1.8 × 1.0 m (brend panel, A4 akril cho'ntaklar — faxriylar, oy o'quvchisi nomzodlari)", 250, "$/dona", "TAXMIN"),
    ('resepshn', "Resepshn stoykasi 1.8 m (buyurtma: peshtaxta 1.10 m, zal tomonidagi panelda logotip uchun joy)", 1500, "$/dona", "TAXMIN — buyurtma asosida yasaladi"),
    ('resepshnPanel', "Resepshn to'siq panellari (past, L shaklida ~3.4 m) va xodim eshikchasi", 500, "$/komplekt", "TAXMIN"),
    ('logoDevor1', "1-qavat: devordagi yorituvchi logotip ~1.5 m (to'q panelda)", 500, "$/dona", "TAXMIN"),
    ('logoStol1', "1-qavat: resepshn stoykasidagi logotip ~1.0 m (yoritilgan)", 250, "$/dona", "TAXMIN"),
    ('turniketSoni', "Turniketlar (buyurtmachida mavjud: kirish va chiqish)", 2, "dona", "Buyurtmachi: mavjud — narxi hisobga kirmaydi"),
    ('zaxira', "Kutilmagan xarajatlar", 0.10, "ulush", "Buyurtmachi: umumiy summadan 10%"),
    ('H', "Devor balandligi (poldan shiftgacha)", 3.5, "m", "Buyurtmachi: bino balandligi 3.5 m"),
    ('Hsh', "Shisha qism balandligi", 2.5, "m", "Buyurtmachi"),
    ('Hust', "Shisha ustidagi GKL", '=C{H}-C{Hsh}', "m", "Hisob: devor balandligi − shisha"),
    ('Hesh', "Eshik balandligi", 2.1, "m", "Taxmin: standart eshik 900 × 2100"),
    ('HeUst', "Eshik ustidagi GKL", '=C{H}-C{Hesh}', "m", "Hisob: devor balandligi − eshik"),
    ('eshEn', "Eshik eni", 0.9, "m", "Chizmadan"),
    ('Hb', "Bambuk panel balandligi", '=C{H}', "m", "Taxmin: poldan shiftgacha. Pastroq bo'lsa, shu katakka raqam yozing"),
    ('HbYQ', "Bambuk panel balandligi — yuqori qavat o'quv xonasi", 2.5, "m", "Buyurtmachi"),
]

# ---------- Batafsil qatorlar ----------
qatorlar = {k: [] for k, _ in XONALAR}


def qosh(xona, tur, joy, uz=None, bal=None, miqdor=None, birlik='m²', narx=None, manfiy=False):
    # manfiy=True — ayiriladigan maydon (masalan, eshik o'rni)
    qatorlar[xona].append(dict(tur=tur, joy=joy, uz=uz, bal=bal, miqdor=miqdor, birlik=birlik, narx=narx, manfiy=manfiy))


def nomlar(kodlar):
    return ', '.join(NOM.get(k, k) for k in kodlar)


def qarshi(f, yuz):
    # devorning qarshi tomonidagi, shu yuz bilan ustma-ust tushadigan xonalar
    return [g['kod'] for g in yuz if g['tomon'] != f['tomon'] and min(f['a2'], g['a2']) > max(f['a1'], g['a1'])]


for w in d['devorlar']:
    yuz = [f for f in w['yuz'] if f['kod'] in NOM]
    if not yuz:
        continue  # 10 sm lik ustuncha (shisha devorlar orasida) — hisobga kirmaydi
    if w['shisha']:
        egasi = next(f for f in yuz if f['kod'] not in UMUMIY)
        kor = [f['kod'] for f in yuz if f['kod'] in UMUMIY]
        qosh(egasi['kod'], 'Shisha devor', f"Shisha devor ({nomlar(kor)} tomoni)", egasi['l'] / 1000, 'Hsh', narx='shisha')
        for f in yuz:
            qosh(f['kod'], 'Gips karton (GKL)', f"Shisha ustidagi GKL ({nomlar(qarshi(f, yuz))} bilan orasida)", f['l'] / 1000, 'Hust', narx='gkl')
        continue
    for f in yuz:
        boshqa = qarshi(f, yuz)
        if w.get('logo'):
            joy = (f"Koridor oxiridagi logotip devori ({nomlar(boshqa)} bilan orasida)" if f['kod'] in UMUMIY
                   else f"{nomlar(boshqa)} oxiridagi logotip devori (xona tomoni)")
            qosh(f['kod'], 'Gips karton (GKL)', joy, f['l'] / 1000, 'H', narx='gkl')
            continue
        if not boshqa:
            joy = "Koridor uchidagi burchak ustunchasi"
        elif f['l'] <= 550 and not w['izoh']:
            joy = f"Eshik yonidagi kesak ({nomlar(boshqa)} tomoni)"
        else:
            joy = f"{nomlar(boshqa)} bilan orasidagi devor" + (" — mavjud devordagi bo'shliqni yopish" if w['izoh'] else '')
        qosh(f['kod'], 'Gips karton (GKL)', joy, f['l'] / 1000, 'H', narx='gkl')

# Yangi eshiklar shisha: eshik (2.1 m) va ustidagi shisha (0.4 m) — shisha devor bilan bir chiziqda 2.5 m gacha;
# ustida, shisha devordagi kabi, 1 m GKL ikki tomondan.
for e in d['eshiklar']:
    kor = KORIDOR.get(e['kod']) or next(s['koridor'] for s in d['sinflar'] if s['kod'] == e['kod'])
    qosh(e['kod'], 'Shisha devor', f"Shisha eshik va ustidagi shisha ({NOM[kor]} tomoni)", 'eshEn', 'Hsh', narx='shisha')
    qosh(e['kod'], 'Gips karton (GKL)', f"Eshik ustidagi GKL ({NOM[kor]} bilan orasida)", 'eshEn', 'Hust', narx='gkl')
    qosh(kor, 'Gips karton (GKL)', f"Eshik ustidagi GKL ({NOM[e['kod']]} eshigi)", 'eshEn', 'Hust', narx='gkl')

for s in d['sinflar']:
    qosh(s['kod'], 'Bambuk panel', "Doska devori", s['doskaDevor'] / 1000, 'Hb', narx='bambuk')
    qosh(s['kod'], 'Bambuk panel', "Orqa devor" + (" (L shaklining keng qismi)" if s['kod'] == 'SR4' else ''), s['orqaDevor'] / 1000, 'Hb', narx='bambuk')
    qosh(s['kod'], 'Interaktiv doska', "Doska devoriga", miqdor=1, birlik='dona', narx='doska')
for o in d['ofislar']:
    for dv in o['devorlar']:
        qosh(o['kod'], 'Bambuk panel', dv['nom'], dv['uz'] / 1000, 'Hb', narx='bambuk')
        if dv.get('eshik'):
            qosh(o['kod'], 'Bambuk panel', "Eshik o'rni (ayiriladi)", 'eshEn', 'Hesh', narx='bambuk', manfiy=True)
# Yuqori qavatdagi o'quv xonasi (buyurtmachi bergan chizma: 5.45 × 8.60 m, 46.87 m²) — 3 devorga 2.5 m balandlikda:
# doska devori va unga qarama-qarshi devor (5.45 m), eshikli uzun devor (8.60 m); derazali uzun devorga qilinmaydi.
qosh('YQ', 'Bambuk panel', "Doska devori (eni)", 5.45, 'HbYQ', narx='bambuk')
qosh('YQ', 'Bambuk panel', "Orqa devor (eni)", 5.45, 'HbYQ', narx='bambuk')
qosh('YQ', 'Bambuk panel', "Uzun devor (eshikli)", 8.60, 'HbYQ', narx='bambuk')
qosh('YQ', 'Bambuk panel', "Eshik o'rni (ayiriladi)", 'eshEn', 'Hesh', narx='bambuk', manfiy=True)
qosh('YQ', 'Interaktiv doska', "Shimoliy devorga (eshikdan kirganda chapda)", miqdor=1, birlik='dona', narx='doska')
# Hojatxonalar: santexnika ikkala qavat uchun jami (buyurtmachi bergan sonlar), kafel — faqat erkaklar hojatxonasi
for kalit, nom in [('chashagen', "Chashagen — 3 ta erkaklar (3-qavat), 1 ta ayollar (2-qavat)"),
                   ('unitaz', "Unitaz — 1 ta erkaklar, 2 ta ayollar"),
                   ('rakovina', "Rakovina — 2 ta erkaklar, 2 ta ayollar"),
                   ('mustahab', "Mustahab — har bir unitaz va chashagenga"),
                   ('taxoratSm', "Taxorat smesiteli — 3 ta erkaklar taxorat joyida, 1 ta ayollar"),
                   ('rakovinaSm', "Rakovina smesiteli — har bir rakovinaga")]:
    qosh('HJ', 'Santexnika', nom, miqdor=f'{kalit}Soni', birlik='dona', narx=('som', kalit))
qosh('HJ', 'Santexnika', "Santexnika nuqtalari — usta ishi (suv va kanalizatsiya)", miqdor='tochkaSoni', birlik='nuqta', narx=('som', 'tochka'))
qosh('HJ', 'Santexnika', "Hojatxona kabinalari (erkaklar, 3-qavat)", miqdor='kabinaSoni', birlik='dona', narx=('som', 'kabina'))
qosh('HJ', 'Santexnika', "Issiq suv quvuri", miqdor='issiqUz', birlik='m', narx=('som', 'issiqQuvur'))
qosh('HJ', 'Santexnika', "Sovuq suv quvuri", miqdor='sovuqUz', birlik='m', narx=('som', 'sovuqQuvur'))
qosh('HJ', 'Santexnika', "Quvur aksessuarlari (fitinglar)", miqdor='aksessuarSoni', birlik='dona', narx=('som', 'aksessuar'))
qosh('HJ', 'Santexnika', "Kanalizatsiya quvuri (40 m) va otvodlar (jami)", miqdor=1, birlik='komplekt', narx=('som', 'kanal'))
qosh('HJ', 'Kafel', "Erkaklar hojatxonasi devori — kafel (material)", miqdor='kafelDevor', birlik='m²', narx=('som', 'kafel'))
qosh('HJ', 'Kafel', "Erkaklar hojatxonasi devori — yotqizish (usta)", miqdor='kafelDevor', birlik='m²', narx=('som', 'kafelUsta'))
qosh('HJ', 'Kafel', "Erkaklar hojatxonasi poli — kafel (material)", miqdor='kafelPol', birlik='m²', narx=('som', 'kafel'))
qosh('HJ', 'Kafel', "Erkaklar hojatxonasi poli — yotqizish (usta)", miqdor='kafelPol', birlik='m²', narx=('som', 'kafelUsta'))
qosh('ZINA', 'Perila', "Perila (buyurtmachi bergan uzunlik)", 'perilaUz', None, birlik='m', narx='perila')

# Yoritish (buyurtmachi narxi joy boshiga): koridorlar va koworking — katta, xonalar — kichik komplekt
for kod in ['SR1', 'SR2', 'SR3', 'SR4', 'SR5', 'SR6', 'SR7', 'XZ1', 'XZ2', 'XZ3', 'XZ4', 'X12']:
    qosh(kod, 'Yoritish', "Lampalar (xona uchun)", miqdor=1, birlik='komplekt', narx='yorXona')
for kod in ['K1', 'K2', 'KW']:
    qosh(kod, 'Yoritish', "Lampalar", miqdor=1, birlik='komplekt', narx='yorKatta')
# Videokuzatuv va logotiplar — butun qavat uchun
qosh('BQ', 'Kamera tizimi', "Videokamera", miqdor='kameraSoni', birlik='dona', narx='kamera')
qosh('BQ', 'Kamera tizimi', "Kamera o'rnatish", miqdor='kameraSoni', birlik='dona', narx='kameraOrn')
qosh('BQ', 'Kamera tizimi', "NVR (videoregistrator)", miqdor=1, birlik='dona', narx='nvr')
qosh('BQ', 'Kamera tizimi', "Kabellar (jami)", miqdor=1, birlik='komplekt', narx='kabel')
qosh('BQ', 'Logotip', "Logotiplar — 2-qavat, jami (koridor oxiridagi 2 ta yorituvchi logotip bilan)", miqdor=1, birlik='komplekt', narx='logo')
# So'mdagi ishlar (narx $ = so'm ÷ kurs)
qosh('BQ', 'Matli plyonka', "Shisha devorlarga matli plyonka", miqdor='plyonkaSoni', birlik='dona', narx=('som', 'plyonka'))
qosh('BQ', "Yong'in xavfsizligi", "Signalizatsiya, o't o'chirgichlar, evakuatsiya belgilari, avariya chiroqlari (jami)", miqdor=1, birlik='komplekt', narx=('som', 'yongin'))
for w in d['buziladi']:
    joy = f"{w['izoh']} o'rnidagi bo'lak" if 'koridori' in w['izoh'] else w['izoh']
    qosh('BQ', 'Buzish ishlari', f"Devorni buzish — {joy}", w['uz'] / 1000, 'H', narx=('som', 'buzish'))
qosh('BQ', 'Buzish ishlari', "Qurilish chiqindisini olib chiqish (jami)", miqdor=1, birlik='komplekt', narx=('som', 'chiqindi'))
# Konditsioner va ventilyatsiya (buyurtmachi summalari, o'rnatish bilan)
qosh('BQ', 'Konditsioner va ventilyatsiya', "Konditsioner — 2 ta, 100 mingtalik, o'rnatish bilan (jami)", miqdor=1, birlik='komplekt', narx='kond2')
qosh('BQ', 'Konditsioner va ventilyatsiya', "Havo kanallari (vozduxovod), 200 m² (jami)", miqdor=1, birlik='komplekt', narx='vozdux')
qosh('YQ', 'Konditsioner va ventilyatsiya', "Konditsioner (o'quv xonasi)", miqdor=1, birlik='dona', narx='kond3')

# Sotuv (XZ1) va admin (XZ2) xonalari: mebel va interyer — taxminiy (jihozlar soni chizmadagi kabi)
qosh('XZ1', 'Mebel va interyer', "Konsultant stoli", miqdor=2, birlik='dona', narx='stolOfis')
qosh('XZ1', 'Mebel va interyer', "Konsultant kreslosi", miqdor=2, birlik='dona', narx='kreslo')
qosh('XZ1', 'Mebel va interyer', "Mijoz stuli (har konsultantga 2 ta)", miqdor=4, birlik='dona', narx='mijozStul')
qosh('XZ1', 'Mebel va interyer', "Kutish divani", miqdor=1, birlik='dona', narx='divan')
qosh('XZ1', 'Mebel va interyer', "Hujjat shkafi", miqdor=1, birlik='dona', narx='shkaf')
qosh('XZ1', 'Mebel va interyer', "Interyer: brend devor, o'simliklar, dekor", miqdor=1, birlik='komplekt', narx='interyerXZ1')
qosh('XZ2', 'Mebel va interyer', "Ish stoli", miqdor=2, birlik='dona', narx='stolOfis')
qosh('XZ2', 'Mebel va interyer', "Ofis kreslosi", miqdor=2, birlik='dona', narx='kreslo')
qosh('XZ2', 'Mebel va interyer', "Mehmon stuli", miqdor=2, birlik='dona', narx='mijozStul')
qosh('XZ2', 'Mebel va interyer', "Hujjat shkafi", miqdor=1, birlik='dona', narx='shkaf')
qosh('XZ2', 'Mebel va interyer', "Interyer: dekor, o'simliklar", miqdor=1, birlik='komplekt', narx='interyerXZ2')
# Koridorlarning kar devorlarida doskalar (har bir koridorda bittadan) — taxminiy
for kor, xonalar_ in (('K1', "sotuv va admin xonalari orqa devorida"), ('K2', "ustozlar va call-markaz orqa devorida")):
    qosh(kor, "E'lon va e'tirof doskalari", f"E'lonlar doskasi ({xonalar_})", miqdor=1, birlik='dona', narx='elonDoska')
    qosh(kor, "E'lon va e'tirof doskalari", f"E'tirof doskasi — faxriylar, oy o'quvchisi nomzodlari ({xonalar_})", miqdor=1, birlik='dona', narx='etirofDoska')
# 1-qavat: resepshn, mehmonlar joyi, logotiplar — taxminiy; turniket buyurtmachida mavjud
qosh('Q1', 'Mebel va interyer', "Resepshn stoykasi (buyurtma asosida)", miqdor=1, birlik='dona', narx='resepshn')
qosh('Q1', 'Mebel va interyer', "Resepshn to'siq panellari va xodim eshikchasi (L shaklida)", miqdor=1, birlik='komplekt', narx='resepshnPanel')
qosh('Q1', 'Mebel va interyer', "Resepshn xodimi kreslosi", miqdor=1, birlik='dona', narx='kreslo')
qosh('Q1', 'Mebel va interyer', "Mehmonlar divani", miqdor=1, birlik='dona', narx='divan')
qosh('Q1', 'Mebel va interyer', "Jurnal stolchasi", miqdor=1, birlik='dona', narx='stolcha')
qosh('Q1', 'Logotip', "Resepshn orqasidagi o'ng devorda yorituvchi logotip", miqdor=1, birlik='dona', narx='logoDevor1')
qosh('Q1', 'Logotip', "Resepshn stoykasining zal tomonidagi panelidagi logotip", miqdor=1, birlik='dona', narx='logoStol1')

TURLAR = ['Shisha devor', 'Gips karton (GKL)', 'Bambuk panel', 'Interaktiv doska', 'Perila', 'Yoritish', 'Kamera tizimi', 'Logotip',
          'Matli plyonka', "Yong'in xavfsizligi", 'Buzish ishlari', 'Santexnika', 'Kafel', 'Konditsioner va ventilyatsiya',
          'Mebel va interyer', "E'lon va e'tirof doskalari"]
JUFT = TURLAR[:5]    # «Xonalar bo'yicha» varag'ida miqdor va summa ustunlari bilan
YAKKA = TURLAR[5:]   # faqat summa ustuni bilan
for k in qatorlar:
    qatorlar[k].sort(key=lambda q: TURLAR.index(q['tur']))

# ---------- Uslublar ----------
SHRIFT = 'Arial'
f_oddiy = Font(name=SHRIFT, size=10)
f_kok = Font(name=SHRIFT, size=10, color='0000FF')       # qo'lda kiritilgan qiymat
f_yashil = Font(name=SHRIFT, size=10, color='008000')    # boshqa varaqqa havola
f_qalin = Font(name=SHRIFT, size=10, bold=True)
f_sarlavha = Font(name=SHRIFT, size=14, bold=True, color='1F3657')
f_kichik = Font(name=SHRIFT, size=9, italic=True, color='555555')
f_bosh = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
fill_bosh = PatternFill('solid', fgColor='1F3657')
fill_guruh = PatternFill('solid', fgColor='E8EEF6')
fill_jami = PatternFill('solid', fgColor='F2F2F2')
fill_sariq = PatternFill('solid', fgColor='FFFF00')
fill_taxmin = PatternFill('solid', fgColor='FFC000')     # taxminiy narx
ingichka = Side(style='thin', color='BFBFBF')
chegara = Border(left=ingichka, right=ingichka, top=ingichka, bottom=ingichka)
USD = '$#,##0.00;($#,##0.00);"-"'
USD0 = '$#,##0;($#,##0);"-"'
SON = '#,##0.00;(#,##0.00);"-"'
DONA = '#,##0;(#,##0);"-"'

wb = Workbook()

# ---------- 1. Narxlar ----------
wn = wb.active
wn.title = 'Narxlar'
wn['A1'] = "PDP Academy · Xadra filiali (1, 2 va 3-qavat): narxlar va o'lchamlar"
wn['A1'].font = f_sarlavha
wn['A2'] = "Sariq kataklardagi narx va balandliklarni o'zgartirsangiz, «Hisob (batafsil)» va «Xonalar bo'yicha» varaqlari avtomatik qayta hisoblanadi."
wn['A2'].font = f_kichik
for i, h in enumerate(['№', "Ko'rsatkich", 'Qiymat', 'Birlik', 'Manba / izoh'], 1):
    c = wn.cell(row=4, column=i, value=h)
    c.font, c.fill, c.border = f_bosh, fill_bosh, chegara
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
r = 5
for i, (k, _, _, _, _) in enumerate(NARX_QATORLAR):
    N[k] = r + i
for i, (k, nom, qiymat, birlik, manba) in enumerate(NARX_QATORLAR):
    rr = r + i
    if isinstance(qiymat, str):
        qiymat = qiymat.format(**N)
    vals = [i + 1, nom, qiymat, birlik, manba]
    for j, v in enumerate(vals, 1):
        c = wn.cell(row=rr, column=j, value=v)
        c.font, c.border = f_oddiy, chegara
    q = wn.cell(row=rr, column=3)
    if isinstance(qiymat, str):
        q.font = f_oddiy
        if k == 'Hb':
            q.fill = fill_sariq   # sukut bo'yicha devor balandligi, lekin o'zgartirish mumkin
    else:
        q.font, q.fill = f_kok, (fill_taxmin if manba.startswith('TAXMIN') else fill_sariq)
    q.number_format = (USD if birlik.startswith('$') else '#,##0 "so\'m"' if birlik.startswith("so'm/") else '#,##0' if birlik == "so'm/$"
                       else '0%' if birlik == 'ulush' else '0' if birlik in ('dona', 'nuqta') else '0.00')
    wn.cell(row=rr, column=1).alignment = Alignment(horizontal='center')
oxir = r + len(NARX_QATORLAR)
izohlar = [
    "Hisoblash qoidalari",
    "• Gips karton (GKL) devor: uzunlik × balandlik. Devorning har bir tomoni o'zi qaragan xonaga yoziladi (xona tomoni — xonaga, koridor tomoni — koridorga).",
    "• Shisha devor: shisha qism (2.5 m) xonaga yoziladi; ustidagi GKL (1 m) ikki tomonlama — xona va koridor tomoni alohida.",
    "• Eshiklar shisha: eshik (2.1 m) va ustidagi shisha (0.4 m) — 2.5 m gacha shisha narxida; ustida 1 m GKL ikki tomondan. Eshik furniturasi (ilgak, qulf, tutqich, dovodchik) kirmagan.",
    "• Bambuk panel, poldan shiftgacha: o'quv xonalarida doska devori va orqa devor (interaktiv doska orqasi ayirilmagan); ofis xonalarida ikkala uzun yon devor; CEO xonasida ikkita ichki kar devor, eshik o'rni ayiriladi. Derazali tashqi va shisha devorlarga qilinmaydi.",
    "• Yuqori qavat o'quv xonasi (5.45 × 8.60 m, buyurtmachi chizmasi): bambuk panel 3 devorga, 2.5 m balandlikda — doska devori, orqa devor va eshikli uzun devor (eshik o'rni ayiriladi); derazali uzun devorga qilinmaydi.",
    "• Interaktiv doska: 7 ta shu qavat sinflariga va 1 ta yuqori qavatdagi o'quv xonasiga — jami 8 ta.",
    "• Perila: buyurtmachi bergan uzunlik (24 m), «Zinapoya» qatorida.",
    "• Uzunliklar chizma v1.1 dan (mavjud holat chizmasining fotosi asosida) — ish boshlanishidan oldin joyida o'lchanadi.",
    "• Yoritish: koridorlar (K1, K2) va koworking zaliga 336 $ dan, o'quv xonalari (7), ofis xonalari (4) va CEO xonasiga 168 $ dan. Kirish zali, WC va yuqori qavat xonasi kiritilmagan.",
    "• Videokuzatuv: 16 ta kamera (40 $) va o'rnatish (21 $ dan), NVR, kabellar jami; logotiplar jami — «Butun qavat» qatorida.",
    "• So'mdagi narxlar (matli plyonka, yong'in xavfsizligi, devor buzish va chiqindi) «Dollar kursi» katagi bo'yicha (11 900 so'm/$) $ ga o'tkaziladi.",
    "• Devor buzish: chizmadagi buziladigan devorlar uzunligi × devor balandligi (3.5 m) × 50 000 so'm/m².",
    "• Kutilmagan xarajatlar: umumiy summadan 10% — «Xonalar bo'yicha» varag'ining oxirida.",
    "• 3-qavat: 5-xona — o'quv xonasi (24 o'rin, bambuk panel va interaktiv doska); 2-xona — erkaklar hojatxonasi. Santexnika soni ikkala hojatxona uchun jami (2-qavatdagisi — ayollar: 2 unitaz, 1 chashagen); kafel faqat erkaklar hojatxonasi uchun (devor 58 m², pol 16 m²).",
    "• Konditsioner va ventilyatsiya (buyurtmachi summalari, o'rnatish bilan): 2-qavatga 2 ta 100 mingtalik konditsioner va 200 m² havo kanali — «Butun qavat» qatorida; 3-qavat o'quv xonasiga 1 ta konditsioner.",
    "• Santexnika: 4 ta kabina, issiq va sovuq suv quvurlari (40 m dan), 100 ta aksessuar, kanalizatsiya quvuri va otvodlar — «Hojatxonalar» qatorida.",
    "• 2-qavat koridorlari: 4-xonaning koridor oxiridagi devorlari GKL (avval shisha edi), markazida yorituvchi logotip; e'lonlar va e'tirof doskalari — K1 va K2 qatorlarida.",
    "• TAXMINIY narxlar (to'q sariq fon): sotuv va admin xonalari mebeli va interyeri, koridor doskalari, 1-qavat resepshn stoykasi, mehmonlar mebeli va logotiplari — Toshkent bozori bo'yicha dastlabki baho; mebelchi va reklama ustasi narxi bilan almashtiring. «Yakuniy» varag'ida taxminiy qism alohida ko'rsatilgan.",
    "• Kirmagan: o'quv xonalari, ustozlar xonasi, call-markaz, CEO va koworking mebeli — mavjud; pol (kafel) va internet — mavjud; turniketlar (2 ta) — mavjud (o'rnatish, to'siq va evakuatsiya darvozasi alohida aniqlanadi); eshik furniturasi; shift va bo'yoq; ish haqi (agar narxga kirmagan bo'lsa).",
    "Ranglar: ko'k — qo'lda kiritilgan qiymat; yashil — boshqa varaqdan olingan; qora — formula; sariq fon — o'zgartirish mumkin bo'lgan kataklar.",
]
for i, t in enumerate(izohlar):
    c = wn.cell(row=oxir + 1 + i, column=2, value=t)
    c.font = f_qalin if i == 0 else f_kichik
for col, w in zip('ABCDE', [5, 38, 12, 9, 80]):
    wn.column_dimensions[col].width = w


def narx_havola(k):
    if isinstance(k, tuple):   # ('som', kalit) — so'mdagi narx, kurs bo'yicha $ ga
        return f"=Narxlar!$C${N[k[1]]}/Narxlar!$C${N['kurs']}"
    return f"=Narxlar!$C${N[k]}"


# ---------- 2. Hisob (batafsil) ----------
wd = wb.create_sheet('Hisob (batafsil)')
wd['A1'] = "PDP Academy · Xadra filiali (1, 2 va 3-qavat): xarajatlar hisobi, xonalar bo'yicha batafsil"
wd['A1'].font = f_sarlavha
wd['A2'] = "Uzunliklar chizmadan (m). Balandlik va narx «Narxlar» varag'idan olinadi (yashil). Miqdor = uzunlik × balandlik; summa = miqdor × narx."
wd['A2'].font = f_kichik
SARLAVHA = ['№', 'Xona', 'Ish turi', 'Joyi / izoh', 'Uzunlik, m', 'Balandlik, m', 'Miqdor', 'Birlik', 'Narx, $', 'Summa, $']
for i, h in enumerate(SARLAVHA, 1):
    c = wd.cell(row=4, column=i, value=h)
    c.font, c.fill, c.border = f_bosh, fill_bosh, chegara
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
r = 5
n = 0
BIRINCHI = r
xona_jami = {}
for kod, nom in XONALAR:
    if not qatorlar[kod]:
        continue
    c = wd.cell(row=r, column=1, value=nom)
    c.font = f_qalin
    for j in range(1, 11):
        wd.cell(row=r, column=j).fill = fill_guruh
    r += 1
    bosh = r
    for q in qatorlar[kod]:
        n += 1
        uz = q['uz']
        uz_val = narx_havola(uz) if isinstance(uz, str) else uz
        bal_val = narx_havola(q['bal']) if q['bal'] else None
        if q['miqdor'] is not None:
            miqdor = narx_havola(q['miqdor']) if isinstance(q['miqdor'], str) else q['miqdor']
        elif q['birlik'] == 'm²':
            miqdor = f"=-E{r}*F{r}" if q['manfiy'] else f"=E{r}*F{r}"
        else:
            miqdor = f"=E{r}"
        vals = [n, nom, q['tur'], q['joy'], uz_val, bal_val, miqdor, q['birlik'], narx_havola(q['narx']), f"=G{r}*I{r}"]
        for j, v in enumerate(vals, 1):
            c = wd.cell(row=r, column=j, value=v)
            c.font, c.border = f_oddiy, chegara
        wd.cell(row=r, column=5).font = f_yashil if isinstance(uz, str) else f_kok
        if bal_val:
            wd.cell(row=r, column=6).font = f_yashil
        if q['miqdor'] is not None:
            wd.cell(row=r, column=7).font = f_yashil if isinstance(q['miqdor'], str) else f_kok
        wd.cell(row=r, column=9).font = f_yashil
        wd.cell(row=r, column=1).alignment = Alignment(horizontal='center')
        wd.cell(row=r, column=8).alignment = Alignment(horizontal='center')
        for j, fmt in [(5, SON), (6, SON), (7, SON if q['birlik'] in ('m²', 'm') else DONA), (9, USD), (10, USD)]:
            wd.cell(row=r, column=j).number_format = fmt
        r += 1
    c = wd.cell(row=r, column=4, value=f"Jami: {nom}")
    c.font = f_qalin
    c = wd.cell(row=r, column=10, value=f"=SUM(J{bosh}:J{r - 1})")
    c.font, c.number_format = f_qalin, USD
    for j in range(1, 11):
        wd.cell(row=r, column=j).fill = fill_jami
        wd.cell(row=r, column=j).border = chegara
    xona_jami[kod] = r
    r += 2
OXIRGI = r - 1
c = wd.cell(row=r, column=4, value="JAMI (barcha xonalar)")
c.font = f_qalin
c = wd.cell(row=r, column=10, value='=' + '+'.join(f"J{x}" for x in xona_jami.values()))
c.font, c.number_format = f_qalin, USD
for j in range(1, 11):
    wd.cell(row=r, column=j).fill = fill_jami
    wd.cell(row=r, column=j).border = chegara
for col, w in zip('ABCDEFGHIJ', [5, 24, 18, 58, 11, 11, 10, 7, 11, 13]):
    wd.column_dimensions[col].width = w
wd.freeze_panes = 'A5'

# ---------- 3. Xonalar bo'yicha ----------
ws = wb.create_sheet("Xonalar bo'yicha", 0)
ws['A1'] = "PDP Academy · Xadra filiali (1, 2 va 3-qavat): xarajatlar, xonalar bo'yicha"
ws['A1'].font = f_sarlavha
ws['A2'] = "Har bir qator «Hisob (batafsil)» varag'idan olinadi (SUMIFS). Narxlarni «Narxlar» varag'ida o'zgartiring."
ws['A2'].font = f_kichik
UST = ['Xona', 'Shisha (devor va eshik), m²', 'Shisha (devor va eshik), $', 'Gips karton (GKL), m²', 'Gips karton (GKL), $',
       'Bambuk panel, m²', 'Bambuk panel, $', 'Interaktiv doska, dona', 'Interaktiv doska, $', 'Perila, m', 'Perila, $',
       'Yoritish, $', 'Kamera tizimi, $', 'Logotip, $', 'Matli plyonka, $', "Yong'in xavfsizligi, $", 'Buzish ishlari, $',
       'Santexnika, $', 'Kafel, $', 'Konditsioner va ventilyatsiya, $', 'Mebel va interyer (taxmin), $', "E'lon va e'tirof doskalari (taxmin), $", 'Jami, $']
NC = len(UST)                      # ustunlar soni
JAMI_UST = get_column_letter(NC)   # «Jami, $» ustuni
for i, h in enumerate(UST, 1):
    c = ws.cell(row=4, column=i, value=h)
    c.font, c.fill, c.border = f_bosh, fill_bosh, chegara
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
ws.row_dimensions[4].height = 32
D_ = "'Hisob (batafsil)'"
rng = lambda col: f"{D_}!${col}${BIRINCHI}:${col}${OXIRGI}"
r = 5
bolim_jami = []
for bolim, kodlar in BOLIM:
    c = ws.cell(row=r, column=1, value=bolim)
    c.font = f_qalin
    for j in range(1, NC + 1):
        ws.cell(row=r, column=j).fill = fill_guruh
    r += 1
    bosh = r
    for kod in kodlar:
        nom = NOM[kod]
        ws.cell(row=r, column=1, value=nom).font = f_oddiy
        for t, tur in enumerate(JUFT):
            mq = ws.cell(row=r, column=2 + 2 * t, value=f'=SUMIFS({rng("G")},{rng("B")},$A{r},{rng("C")},"{tur}")')
            sm = ws.cell(row=r, column=3 + 2 * t, value=f'=SUMIFS({rng("J")},{rng("B")},$A{r},{rng("C")},"{tur}")')
            mq.font = sm.font = f_yashil
            mq.number_format = DONA if tur == 'Interaktiv doska' else SON
            sm.number_format = USD0
        for t, tur in enumerate(YAKKA):
            sm = ws.cell(row=r, column=2 + 2 * len(JUFT) + t, value=f'=SUMIFS({rng("J")},{rng("B")},$A{r},{rng("C")},"{tur}")')
            sm.font, sm.number_format = f_yashil, USD0
        summa_ust = [get_column_letter(3 + 2 * t) for t in range(len(JUFT))] + [get_column_letter(2 + 2 * len(JUFT) + t) for t in range(len(YAKKA))]
        jm = ws.cell(row=r, column=NC, value='=' + '+'.join(f"{L}{r}" for L in summa_ust))
        jm.font, jm.number_format = f_qalin, USD0
        for j in range(1, NC + 1):
            ws.cell(row=r, column=j).border = chegara
        r += 1
    ws.cell(row=r, column=1, value=f"{bolim} — jami").font = f_qalin
    for j in range(2, NC + 1):
        L = get_column_letter(j)
        c = ws.cell(row=r, column=j, value=f"=SUM({L}{bosh}:{L}{r - 1})")
        c.font = f_qalin
        c.number_format = ws.cell(row=r - 1, column=j).number_format
    for j in range(1, NC + 1):
        ws.cell(row=r, column=j).fill = fill_jami
        ws.cell(row=r, column=j).border = chegara
    bolim_jami.append(r)
    r += 2
ws.cell(row=r, column=1, value="JAMI").font = Font(name=SHRIFT, size=11, bold=True)
for j in range(2, NC + 1):
    L = get_column_letter(j)
    c = ws.cell(row=r, column=j, value='=' + '+'.join(f"{L}{x}" for x in bolim_jami))
    c.font = Font(name=SHRIFT, size=11, bold=True)
    c.number_format = ws.cell(row=bolim_jami[0], column=j).number_format
for j in range(1, NC + 1):
    ws.cell(row=r, column=j).fill = PatternFill('solid', fgColor='D9E1F2')
    ws.cell(row=r, column=j).border = chegara
JAMI_QATOR = r
# kutilmagan xarajatlar va umumiy jami ($ va so'mda)
yakun = [
    (f'="Kutilmagan xarajatlar ("&TEXT(Narxlar!$C${N["zaxira"]},"0%")&" — «Narxlar» varag\'idan)"', f"={JAMI_UST}{JAMI_QATOR}*Narxlar!$C${N['zaxira']}", USD0, False),
    ("UMUMIY JAMI (kutilmagan xarajatlar bilan), $", f"={JAMI_UST}{JAMI_QATOR}+{JAMI_UST}{JAMI_QATOR + 1}", USD0, True),
    ("UMUMIY JAMI, so'm (kurs bo'yicha)", f"={JAMI_UST}{JAMI_QATOR + 2}*Narxlar!$C${N['kurs']}", '#,##0 "so\'m"', True),
]
for i, (nom, formula, fmt, qalin) in enumerate(yakun):
    rr = JAMI_QATOR + 1 + i
    ws.cell(row=rr, column=1, value=nom).font = Font(name=SHRIFT, size=11 if qalin else 10, bold=qalin)
    c = ws.cell(row=rr, column=NC, value=formula)
    c.font, c.number_format = Font(name=SHRIFT, size=11 if qalin else 10, bold=qalin), fmt
    for j in range(1, NC + 1):
        ws.cell(row=rr, column=j).fill = PatternFill('solid', fgColor='D9E1F2' if qalin else 'F2F2F2')
        ws.cell(row=rr, column=j).border = chegara
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=NC - 1)
r = JAMI_QATOR + len(yakun) + 2
ws.cell(row=r, column=1, value="Tekshiruv: batafsil varaqdagi jami bilan farq (0 bo'lishi kerak)").font = f_kichik
c = ws.cell(row=r, column=NC, value=f"={JAMI_UST}{JAMI_QATOR}-{D_}!J{OXIRGI + 1}")
c.font, c.number_format = f_kichik, USD
ws.column_dimensions['A'].width = 28
for j in range(2, NC + 1):
    ws.column_dimensions[get_column_letter(j)].width = 12
ws.column_dimensions[JAMI_UST].width = 14
ws.freeze_panes = 'B5'

# ---------- 0. Yakuniy: xarajat turlari va bo'limlar bo'yicha jami ----------
wy = wb.create_sheet('Yakuniy', 0)
wy['A1'] = "PDP Academy · Xadra filiali (1, 2 va 3-qavat): jami xarajatlar — yakuniy"
wy['A1'].font = f_sarlavha
wy['A2'] = "Barcha summalar «Hisob (batafsil)» varag'idan formulalar bilan olinadi; so'm — «Narxlar» varag'idagi kurs bo'yicha."
wy['A2'].font = f_kichik
KURS = f"Narxlar!$C${N['kurs']}"
def sarlavha(row, nomlar):
    for i, h in enumerate(nomlar, 1):
        c = wy.cell(row=row, column=i, value=h)
        c.font, c.fill, c.border = f_bosh, fill_bosh, chegara
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
def qator(row, vals, fmts, font=f_oddiy, fill=None):
    for i, (v, fm) in enumerate(zip(vals, fmts), 1):
        c = wy.cell(row=row, column=i, value=v)
        c.font, c.border = font, chegara
        if fm: c.number_format = fm
        if fill: c.fill = fill
    wy.cell(row=row, column=1).alignment = Alignment(horizontal='center')
SOM = '#,##0 "so\'m"'
r = 4
wy.cell(row=r, column=1, value="1. Xarajat turlari bo'yicha").font = f_qalin
r += 1
sarlavha(r, ['№', 'Xarajat turi', 'Summa, $', "Summa, so'm", 'Ulush'])
r += 1
tur_bosh = r
for i, tur in enumerate(TURLAR, 1):
    qator(r, [i, tur, f'=SUMIFS({rng("J")},{rng("C")},"{tur}")', f"=C{r}*{KURS}", f"=IF($C${tur_bosh + len(TURLAR)}=0,0,C{r}/$C${tur_bosh + len(TURLAR)})"],
          [None, None, USD, SOM, '0.0%'])
    wy.cell(row=r, column=3).font = f_yashil
    r += 1
JY = r
qator(r, ['', 'JAMI', f"=SUM(C{tur_bosh}:C{r - 1})", f"=C{r}*{KURS}", f"=SUM(E{tur_bosh}:E{r - 1})"], [None, None, USD, SOM, '0.0%'], f_qalin, fill_jami)
r += 1
qator(r, ['', f'="Kutilmagan xarajatlar ("&TEXT(Narxlar!$C${N["zaxira"]},"0%")&")"', f"=C{JY}*Narxlar!$C${N['zaxira']}", f"=C{r}*{KURS}", ''], [None, None, USD, SOM, None], f_oddiy, fill_jami)
r += 1
UY = r
qator(r, ['', 'UMUMIY JAMI', f"=C{JY}+C{JY + 1}", f"=C{r}*{KURS}", ''], [None, None, USD, SOM, None], Font(name=SHRIFT, size=11, bold=True), PatternFill('solid', fgColor='D9E1F2'))
r += 1
taxmin_f = (f'=SUMIFS({rng("J")},{rng("C")},"Mebel va interyer")+SUMIFS({rng("J")},{rng("C")},"E\'lon va e\'tirof doskalari")'
            f'+SUMIFS({rng("J")},{rng("B")},"{NOM["Q1"]}",{rng("C")},"Logotip")')
qator(r, ['', "shundan taxminiy narxlar (mebel, interyer, doskalar, 1-qavat logotiplari) — kutilmagan xarajatsiz", taxmin_f, f"=C{r}*{KURS}", ''], [None, None, USD, SOM, None], f_kichik)
wy.cell(row=r, column=3).fill = fill_taxmin
r += 2
wy.cell(row=r, column=1, value="2. Qavat va bo'limlar bo'yicha").font = f_qalin
r += 1
sarlavha(r, ['№', "Bo'lim", 'Summa, $', "Summa, so'm", 'Ulush'])
r += 1
b_bosh = r
SX = "'Xonalar bo''yicha'"
for i, ((bolim, _), qq) in enumerate(zip(BOLIM, bolim_jami), 1):
    qator(r, [i, bolim, f"={SX}!${JAMI_UST}${qq}", f"=C{r}*{KURS}", f"=IF($C${b_bosh + len(BOLIM)}=0,0,C{r}/$C${b_bosh + len(BOLIM)})"], [None, None, USD, SOM, '0.0%'])
    wy.cell(row=r, column=3).font = f_yashil
    r += 1
qator(r, ['', 'JAMI', f"=SUM(C{b_bosh}:C{r - 1})", f"=C{r}*{KURS}", f"=SUM(E{b_bosh}:E{r - 1})"], [None, None, USD, SOM, '0.0%'], f_qalin, fill_jami)
r += 2
wy.cell(row=r, column=2, value="Tekshiruv: ikki jadval jamisi farqi (0 bo'lishi kerak)").font = f_kichik
c = wy.cell(row=r, column=3, value=f"=C{JY}-C{r - 2}")
c.font, c.number_format = f_kichik, USD
r += 1
wy.cell(row=r, column=2, value="Dollar kursi (so'm/$) — «Narxlar» varag'idan").font = f_kichik
c = wy.cell(row=r, column=3, value=f"={KURS}")
c.font, c.number_format = f_yashil, '#,##0'
for col, w in zip('ABCDE', [5, 46, 16, 20, 9]):
    wy.column_dimensions[col].width = w

for sh in wb.worksheets:
    sh.page_setup.orientation = 'landscape'
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.sheet_properties.pageSetUpPr.fitToPage = True

CHIQISH.parent.mkdir(exist_ok=True)
wb.save(CHIQISH)
print(CHIQISH)
